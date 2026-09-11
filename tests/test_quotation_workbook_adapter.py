from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from app.adapters.quotation.workbook import OpenpyxlQuotationWorkbookAdapter
from app.domain.quotation.value_objects import (
    QuotationDetailSheet,
    QuotationFixedChargeRow,
    QuotationSummaryMeta,
    QuotationSummaryRow,
    QuotationWorkbookData,
)


def test_summary_uses_detail_sheet_total_by_sequential_index_and_formula_totals():
    workbook_data = QuotationWorkbookData(
        summary_sheet_name="汇总",
        summary_title="汇总",
        summary_meta=QuotationSummaryMeta(),
        summary_rows=[
            QuotationSummaryRow(
                part_no="A",
                name="外购件",
                quantity_display="1",
                unit_price=Decimal("12"),
                amount=Decimal("12"),
                detail_sheet_name="",
            ),
            QuotationSummaryRow(
                part_no="B",
                name="有明细组件",
                quantity_display="2",
                unit_price=Decimal("99"),
                amount=Decimal("198"),
                detail_sheet_name="组件B",
            ),
        ],
        detail_sheets=[
            QuotationDetailSheet(
                sheet_name="组件B",
                rows=[{"累计用量": 3, "单价": 4, "总价": 12}],
                total_amount=12,
            )
        ],
    )

    export = OpenpyxlQuotationWorkbookAdapter().export_workbook(workbook_data)
    wb = load_workbook(BytesIO(export.content), data_only=False)
    summary_ws = wb["汇总"]

    assert "组件B" in wb.sheetnames
    assert summary_ws["E12"].value == 12
    assert summary_ws["E13"].value == "='组件B'!C3"
    assert summary_ws["F13"].value == "=E13*D13"
    assert summary_ws["F10"].value == "=SUM(F12:F13)"


def test_summary_uses_detail_sheet_price_column_last_row_when_no_total_column():
    workbook_data = QuotationWorkbookData(
        summary_sheet_name="汇总",
        summary_title="汇总",
        summary_meta=QuotationSummaryMeta(),
        summary_rows=[
            QuotationSummaryRow(
                part_no="A",
                name="机架",
                quantity_display="2",
                unit_price=Decimal("88"),
                amount=Decimal("176"),
                detail_sheet_name="机架",
            ),
        ],
        detail_sheets=[
            QuotationDetailSheet(
                sheet_name="机架",
                rows=[
                    {"根父件名称": "机架", "累计用量": 2, "单价": 10, "__root_inv_name": "机架"},
                    {"根父件名称": "机架", "累计用量": "", "单价": 88, "__root_inv_name": "机架"},
                ],
                total_amount=88,
            )
        ],
    )

    export = OpenpyxlQuotationWorkbookAdapter().export_workbook(workbook_data)
    wb = load_workbook(BytesIO(export.content), data_only=False)
    summary_ws = wb["汇总"]

    detail_ws = wb["机架"]

    assert summary_ws["E12"].value == "='机架'!B3"
    assert summary_ws["F12"].value == "=E12*D12"
    assert summary_ws["F10"].value == "=SUM(F12:F12)"
    headers = [cell.value for cell in detail_ws[1]]
    assert "根父件名称" not in headers
    assert "__root_inv_name" not in headers


def test_summary_number_formats_for_quantity_price_and_amount():
    workbook_data = QuotationWorkbookData(
        summary_sheet_name="汇总",
        summary_title="汇总",
        summary_meta=QuotationSummaryMeta(),
        summary_rows=[
            QuotationSummaryRow(
                part_no="A",
                name="外购件",
                quantity_display="1.5",
                unit_price=Decimal("12.5"),
                amount=Decimal("18.75"),
                detail_sheet_name="",
            ),
        ],
    )

    export = OpenpyxlQuotationWorkbookAdapter().export_workbook(workbook_data)
    wb = load_workbook(BytesIO(export.content), data_only=False)
    summary_ws = wb["汇总"]

    # D12 = quantity, E12 = unit_price, F12 = amount formula, F10 = grand total SUM
    assert summary_ws["D12"].number_format == "0.####"
    assert summary_ws["E12"].number_format == "0.00"
    assert summary_ws["F12"].number_format == "0.00"
    assert summary_ws["F10"].number_format == "0.00"
    # Decimal quantity written through display string -> parsed to float 1.5
    assert summary_ws["D12"].value == 1.5


def test_summary_fixed_charge_rows_use_two_decimal_price_and_amount_formats():
    workbook_data = QuotationWorkbookData(
        summary_sheet_name="汇总",
        summary_title="汇总",
        summary_meta=QuotationSummaryMeta(),
        summary_rows=[
            QuotationSummaryRow(
                part_no="A",
                name="外购件",
                quantity_display="1",
                unit_price=Decimal("12"),
                amount=Decimal("12"),
                detail_sheet_name="",
            ),
        ],
        fixed_charge_rows=[QuotationFixedChargeRow(name="杂项")],
    )

    export = OpenpyxlQuotationWorkbookAdapter().export_workbook(workbook_data)
    wb = load_workbook(BytesIO(export.content), data_only=False)
    summary_ws = wb["汇总"]

    # Fixed charge rows start right after summary rows: row 13 here.
    # Layout for fixed charge rows: D=unit_price, E=amount, F=remark
    assert summary_ws["D13"].number_format == "0.00"
    assert summary_ws["E13"].number_format == "0.00"


def test_summary_sequence_number_column_starts_at_one_and_stops_at_labor():
    workbook_data = QuotationWorkbookData(
        summary_sheet_name="汇总",
        summary_title="汇总",
        summary_meta=QuotationSummaryMeta(),
        summary_rows=[
            QuotationSummaryRow(
                part_no="A",
                name="外购件",
                quantity_display="1",
                unit_price=Decimal("12"),
                amount=Decimal("12"),
                detail_sheet_name="",
            ),
            QuotationSummaryRow(
                part_no="B",
                name="组件",
                quantity_display="2",
                unit_price=Decimal("10"),
                amount=Decimal("20"),
                detail_sheet_name="",
            ),
        ],
        fixed_charge_rows=[
            QuotationFixedChargeRow(name="杂项"),
            QuotationFixedChargeRow(name="FOB"),
            QuotationFixedChargeRow(name="运输费"),
            QuotationFixedChargeRow(name="设计工时"),
            QuotationFixedChargeRow(name="人工"),
            QuotationFixedChargeRow(name="额外行"),
        ],
    )

    export = OpenpyxlQuotationWorkbookAdapter().export_workbook(workbook_data)
    wb = load_workbook(BytesIO(export.content), data_only=False)
    summary_ws = wb["汇总"]

    # A11 = 序号 header
    assert summary_ws["A11"].value == "序号"
    # Summary rows: A12=1, A13=2
    assert summary_ws["A12"].value == 1
    assert summary_ws["A13"].value == 2
    # Fixed charge rows continue numbering: 杂项=3, FOB=4, 运输费=5, 设计工时=6, 人工=7
    assert summary_ws["B14"].value == "杂项"
    assert summary_ws["A14"].value == 3
    assert summary_ws["B15"].value == "FOB"
    assert summary_ws["A15"].value == 4
    assert summary_ws["B16"].value == "运输费"
    assert summary_ws["A16"].value == 5
    assert summary_ws["B17"].value == "设计工时"
    assert summary_ws["A17"].value == 6
    assert summary_ws["B18"].value == "人工"
    assert summary_ws["A18"].value == 7
    # Row after 人工 does NOT get a sequence number
    assert summary_ws["B19"].value == "额外行"
    assert summary_ws["A19"].value is None
